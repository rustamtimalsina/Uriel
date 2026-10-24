# backend/app/parser/docx_parser.py
import logging
from typing import Iterator
import docx

logger = logging.getLogger(__name__)

class DOCXParser:
    CHUNK_CHAR_SIZE = 50_000

    def iter_pages(self, path: str) -> Iterator[str]:
        """Yield the docx text in chunks of paragraphs/tables."""
        try:
            doc = docx.Document(path)
        except Exception as e:
            logger.error("Could not parse DOCX %s: %s", path, e)
            return

        buf: list[str] = []
        buf_len = 0
        any_text = False

        try:
            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue
                buf.append(text)
                buf_len += len(text)
                if buf_len >= self.CHUNK_CHAR_SIZE:
                    chunk = "\n".join(buf)
                    any_text = True
                    yield chunk
                    buf, buf_len = [], 0

            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        buf.append(row_text)
                        buf_len += len(row_text)
                        if buf_len >= self.CHUNK_CHAR_SIZE:
                            chunk = "\n".join(buf)
                            any_text = True
                            yield chunk
                            buf, buf_len = [], 0

            if buf:
                chunk = "\n".join(buf)
                if chunk.strip():
                    any_text = True
                yield chunk

            if not any_text:
                logger.warning("Empty DOCX file: %s", path)
        except Exception:
            logger.exception("Unexpected error while streaming DOCX %s", path)

    def parse(self, path: str) -> str | None:
        chunks = list(self.iter_pages(path))
        return "\n\n".join(chunks) if chunks else None
